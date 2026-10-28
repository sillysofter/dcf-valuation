"""Export the model to an Excel workbook with live formulas.

Sheets: Inputs, WACC, DCF (projections, terminal value, equity bridge, treasury stock method) and
Sensitivity. Inputs, WACC and DCF are fully formula-driven, so changing a blue input in Excel flows
through to the share price. Sensitivity tables are values computed in Python (openpyxl cannot write
Excel data tables), so rerun ``dcfmodel`` after changing inputs to refresh them.

Colour convention: blue = hard-coded input, black = formula, green = link to another sheet.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .sensitivity import SensitivityTable
from .valuation import DCFResult

BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BASE_FILL = PatternFill("solid", fgColor="FFF2CC")
TOP_BORDER = Border(top=Side(style="thin"))

NUM = "#,##0.0_);(#,##0.0)"
PCT = "0.0%_);(0.0%)"
PCT2 = "0.00%_);(0.00%)"
MULT = '0.0"x"'
PRICE = "#,##0.00_);(#,##0.00)"
SHARES = "#,##0.000"
FACTOR = "0.0000"
BETA = "0.00"
INT = "0"

CURRENCY_SYMBOLS = {"USD": "$", "EUR": "€", "GBP": "£", "JPY": "¥", "CAD": "C$", "AUD": "A$"}


def _abs(col: int, row: int, sheet: str | None = None) -> str:
    ref = f"${get_column_letter(col)}${row}"
    return f"{sheet}!{ref}" if sheet else ref


def _rel(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _put(ws: Worksheet, col: int, row: int, value, fmt: str | None = None, colour: str = BLACK,
         bold: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(color=colour, bold=bold)
    if fmt:
        cell.number_format = fmt
    return cell


def _title(ws: Worksheet, text: str, subtitle: str) -> None:
    ws["B1"] = text
    ws["B1"].font = Font(bold=True, size=14)
    ws["B2"] = subtitle
    ws["B2"].font = Font(italic=True)
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 44


def _section(ws: Worksheet, row: int, text: str, last_col: int) -> None:
    for c in range(2, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
    ws.cell(row=row, column=2, value=text)


def write_excel(result: DCFResult, tables: list[SensitivityTable], path: str | Path) -> dict[str, str]:
    """Write the workbook and return a map of key outputs to their cell references (e.g. ``"DCF!$C$40"``)."""
    a = result.assumptions
    sym = CURRENCY_SYMBOLS.get(a.currency, "")
    units = f"({sym or a.currency + ' '}in {a.units}, except per-share data)"
    refs: dict[str, str] = {}

    wb = Workbook()
    ws_in = wb.active
    ws_in.title = "Inputs"
    ws_w = wb.create_sheet("WACC")
    ws_d = wb.create_sheet("DCF")
    ws_s = wb.create_sheet("Sensitivity")

    # ---------------------------------------------------------------- Inputs
    _title(ws_in, f"{a.company}: DCF inputs", units)
    ws_in["B3"] = "Blue = hard-coded input, black = formula, green = link to another sheet"
    ws_in["B3"].font = Font(italic=True, color="808080")
    ws_in.column_dimensions["C"].width = 14
    ws_in.column_dimensions["D"].width = 14
    inp: dict[str, str] = {}
    row = 5

    def input_row(key: str, label: str, value, fmt: str) -> None:
        nonlocal row
        ws_in.cell(row=row, column=2, value=label)
        _put(ws_in, 3, row, value, fmt, BLUE)
        inp[key] = _abs(3, row, "Inputs")
        row += 1

    sections = [
        ("General", [
            ("base_year", "Base fiscal year (valuation date = year end)", a.base_year, INT),
            ("mid_year", "Mid-year convention (1 = on, 0 = off)", 1 if a.mid_year else 0, INT),
            ("tax_rate", "Tax rate", a.tax_rate, PCT),
        ]),
        ("Base year", [
            ("base_revenue", "Revenue", a.base_revenue, NUM),
            ("base_nwc", "Net working capital", a.base_nwc, NUM),
        ]),
        ("Terminal value", [
            ("growth", "Terminal (perpetuity) growth rate", a.terminal_growth, PCT),
            ("multiple", "Exit EV/EBITDA multiple", a.exit_multiple, MULT),
        ]),
        ("Equity bridge", [
            ("debt", "Total debt", a.debt, NUM),
            ("preferred", "Preferred equity", a.preferred_equity, NUM),
            ("minority", "Minority (non-controlling) interest", a.minority_interest, NUM),
            ("cash", "Cash and equivalents", a.cash, NUM),
            ("non_op", "Non-operating assets (e.g. investments)", a.non_operating_assets, NUM),
        ]),
        ("Shares", [
            ("basic", "Basic shares outstanding", a.basic_shares, SHARES),
            ("rsus", "RSUs / restricted stock (fully dilutive)", a.rsus, SHARES),
            ("current_price", f"Current share price ({sym or a.currency})", a.current_share_price, PRICE),
        ]),
    ]
    for name, items in sections:
        _section(ws_in, row, name, 4)
        row += 1
        for item in items:
            input_row(*item)
        row += 1

    tranches = sorted((o for o in a.options if o.count > 0), key=lambda o: o.strike)
    _section(ws_in, row, "Options (keep sorted by ascending strike)", 4)
    ws_in.cell(row=row, column=3, value="Count")
    ws_in.cell(row=row, column=4, value="Strike")
    row += 1
    opt_rows = []
    for i, t in enumerate(tranches, start=1):
        ws_in.cell(row=row, column=2, value=f"Tranche {i}")
        _put(ws_in, 3, row, t.count, SHARES, BLUE)
        _put(ws_in, 4, row, t.strike, PRICE, BLUE)
        opt_rows.append(row)
        row += 1
    if not tranches:
        ws_in.cell(row=row, column=2, value="None")

    # ---------------------------------------------------------------- WACC
    _title(ws_w, f"{a.company}: weighted average cost of capital", "CAPM cost of equity, Hamada relevering")
    ws_w.column_dimensions["C"].width = 14
    w: dict[str, str] = {}
    wacc_rows = [
        ("rf", "Risk-free rate", a.risk_free_rate, PCT2, BLUE),
        ("erp", "Equity risk premium", a.equity_risk_premium, PCT2, BLUE),
        ("beta_u", "Unlevered beta", a.unlevered_beta, BETA, BLUE),
        ("size", "Size premium", a.size_premium, PCT2, BLUE),
        ("dc", "Target debt / total capital", a.target_debt_to_capital, PCT, BLUE),
        ("kd", "Pre-tax cost of debt", a.pre_tax_cost_of_debt, PCT2, BLUE),
        ("tax", "Tax rate", f"={inp['tax_rate']}", PCT, GREEN),
        (None, None, None, None, None),
        ("de", "Target debt / equity", "={dc}/(1-{dc})", PCT, BLACK),
        ("beta_l", "Levered beta  = βu × (1 + (1 − t) × D/E)", "={beta_u}*(1+(1-{tax})*{de})", BETA, BLACK),
        ("ke", "Cost of equity  = rf + βL × ERP + size", "={rf}+{beta_l}*{erp}+{size}", PCT2, BLACK),
        ("kd_at", "After-tax cost of debt  = kd × (1 − t)", "={kd}*(1-{tax})", PCT2, BLACK),
        ("we", "Equity weight", "=1-{dc}", PCT, BLACK),
        ("wd", "Debt weight", "={dc}", PCT, BLACK),
        ("wacc", "WACC", "={we}*{ke}+{wd}*{kd_at}", PCT2, BLACK),
    ]
    row = 4
    _section(ws_w, row, "Cost of capital", 3)
    row += 1
    for key, label, value, fmt, colour in wacc_rows:
        if key is None:
            row += 1
            continue
        if isinstance(value, str) and "{" in value:
            value = value.format(**{k: v.split("!")[-1] for k, v in w.items()})
        ws_w.cell(row=row, column=2, value=label)
        _put(ws_w, 3, row, value, fmt, colour, bold=key == "wacc")
        w[key] = _abs(3, row, "WACC")
        row += 1
    refs["wacc"] = w["wacc"]
    refs["levered_beta"] = w["beta_l"]
    refs["cost_of_equity"] = w["ke"]

    # ---------------------------------------------------------------- DCF
    n = a.n_years
    first, last = 4, 4 + n - 1            # projection columns D..; column C is the base year
    _title(ws_d, f"{a.company}: discounted cash flow analysis", units)
    for c in range(3, max(last, 10) + 1):
        ws_d.column_dimensions[get_column_letter(c)].width = 14
    _section(ws_d, 4, "Unlevered free cash flow", last)
    ws_d.cell(row=4, column=3, value=f"FY{a.base_year}A").alignment = Alignment(horizontal="right")
    for t, year in enumerate(a.years):
        ws_d.cell(row=4, column=first + t, value=f"FY{year}E").alignment = Alignment(horizontal="right")

    r = {name: 5 + i for i, name in enumerate([
        "growth", "revenue", "margin", "ebitda", "da_pct", "da", "ebit", "taxes", "nopat", "plus_da",
        "capex_pct", "capex", "nwc_pct", "nwc", "dnwc", "ufcf", "_blank", "period", "disc_period",
        "factor", "pv",
    ])}
    labels = {
        "growth": ("Revenue growth", PCT), "revenue": ("Revenue", NUM), "margin": ("EBITDA margin", PCT),
        "ebitda": ("EBITDA", NUM), "da_pct": ("D&A, % of revenue", PCT), "da": ("Less: D&A", NUM),
        "ebit": ("EBIT", NUM), "taxes": ("Less: taxes on EBIT  = t × max(EBIT, 0)", NUM),
        "nopat": ("NOPAT", NUM), "plus_da": ("Plus: D&A", NUM), "capex_pct": ("CapEx, % of revenue", PCT),
        "capex": ("Less: capital expenditures", NUM), "nwc_pct": ("NWC, % of revenue", PCT),
        "nwc": ("Net working capital", NUM), "dnwc": ("Less: increase in NWC", NUM),
        "ufcf": ("Unlevered free cash flow", NUM), "period": ("Year", INT),
        "disc_period": ("Discount period (years)", "0.0"), "factor": ("Discount factor", FACTOR),
        "pv": ("PV of unlevered free cash flow", NUM),
    }
    for key, (label, _) in labels.items():
        ws_d.cell(row=r[key], column=2, value=label).font = Font(bold=key in ("revenue", "ebitda", "ufcf"))

    _put(ws_d, 3, r["revenue"], f"={inp['base_revenue']}", NUM, GREEN)
    _put(ws_d, 3, r["nwc"], f"={inp['base_nwc']}", NUM, GREEN)
    drivers = {"growth": a.revenue_growth, "margin": a.ebitda_margin, "da_pct": a.da_pct_revenue,
               "capex_pct": a.capex_pct_revenue, "nwc_pct": a.nwc_pct_revenue}
    tax, mid, wacc = inp["tax_rate"], inp["mid_year"], w["wacc"]
    for t in range(n):
        c = first + t
        cur = lambda key: _rel(c, r[key])  # noqa: E731
        prev = lambda key: _rel(c - 1, r[key])  # noqa: E731
        for key, values in drivers.items():
            _put(ws_d, c, r[key], values[t], PCT, BLUE)
        formulas = {
            "revenue": f"={prev('revenue')}*(1+{cur('growth')})",
            "ebitda": f"={cur('revenue')}*{cur('margin')}",
            "da": f"={cur('revenue')}*{cur('da_pct')}",
            "ebit": f"={cur('ebitda')}-{cur('da')}",
            "taxes": f"=MAX({cur('ebit')},0)*{tax}",
            "nopat": f"={cur('ebit')}-{cur('taxes')}",
            "plus_da": f"={cur('da')}",
            "capex": f"={cur('revenue')}*{cur('capex_pct')}",
            "nwc": f"={cur('revenue')}*{cur('nwc_pct')}",
            "dnwc": f"={cur('nwc')}-{prev('nwc')}",
            "ufcf": f"={cur('nopat')}+{cur('plus_da')}-{cur('capex')}-{cur('dnwc')}",
            "period": "=1" if t == 0 else f"={prev('period')}+1",
            "disc_period": f"=IF({mid}=1,{cur('period')}-0.5,{cur('period')})",
            "factor": f"=1/(1+{wacc})^{cur('disc_period')}",
            "pv": f"={cur('ufcf')}*{cur('factor')}",
        }
        for key, formula in formulas.items():
            _put(ws_d, c, r[key], formula, labels[key][1], BLACK,
                 bold=key in ("revenue", "ebitda", "ufcf"))
    for c in range(2, last + 1):
        ws_d.cell(row=r["ufcf"], column=c).border = TOP_BORDER

    refs["ufcf"] = [_abs(first + t, r["ufcf"], "DCF") for t in range(n)]
    refs["revenue"] = [_abs(first + t, r["revenue"], "DCF") for t in range(n)]
    refs["ebitda"] = [_abs(first + t, r["ebitda"], "DCF") for t in range(n)]

    # Terminal value and equity bridge: column C = perpetuity growth, column D = exit multiple.
    P, X = 3, 4
    fcf_n, ebitda_n, n_ref = _abs(last, r["ufcf"]), _abs(last, r["ebitda"]), _abs(last, r["period"])
    pv_range = f"{_abs(first, r['pv'])}:{_abs(last, r['pv'])}"
    row = r["pv"] + 2
    _section(ws_d, row, "Terminal value and enterprise value", 4)
    ws_d.cell(row=row, column=P, value="Perpetuity").alignment = Alignment(horizontal="right")
    ws_d.cell(row=row, column=X, value="Exit multiple").alignment = Alignment(horizontal="right")
    row += 1
    v: dict[str, int] = {}

    def vrow(key: str, label: str, perp, exit_, fmt: str, colours=(BLACK, BLACK), bold=False) -> None:
        nonlocal row
        ws_d.cell(row=row, column=2, value=label).font = Font(bold=bold)
        if perp is not None:
            _put(ws_d, P, row, perp.format(**{k: _rel(P, i) for k, i in v.items()}), fmt, colours[0], bold)
        if exit_ is not None:
            _put(ws_d, X, row, exit_.format(**{k: _rel(X, i) for k, i in v.items()}), fmt, colours[1], bold)
        v[key] = row
        row += 1

    vrow("wacc", "WACC", f"={wacc}", f"={wacc}", PCT2, (GREEN, GREEN))
    vrow("g", "Terminal growth rate", f"={inp['growth']}", None, PCT, (GREEN, GREEN))
    vrow("mult", "Exit EV/EBITDA multiple", None, f"={inp['multiple']}", MULT, (GREEN, GREEN))
    vrow("metric", "Final-year UFCF / final-year EBITDA", f"={fcf_n}", f"={ebitda_n}", NUM)
    vrow("tv", "Terminal value", "={metric}*(1+{g})/({wacc}-{g})", "={metric}*{mult}", NUM)
    vrow("tv_period", "Terminal value discount period",
         f"=IF({mid}=1,{n_ref}-0.5,{n_ref})", f"={n_ref}", "0.0")
    vrow("pv_tv", "PV of terminal value", "={tv}/(1+{wacc})^{tv_period}", "={tv}/(1+{wacc})^{tv_period}", NUM)
    vrow("sum_pv", "Sum of PV of UFCF", f"=SUM({pv_range})", f"=SUM({pv_range})", NUM)
    vrow("ev", "Enterprise value", "={sum_pv}+{pv_tv}", "={sum_pv}+{pv_tv}", NUM, bold=True)
    vrow("tv_pct", "PV of terminal value, % of EV", "={pv_tv}/{ev}", "={pv_tv}/{ev}", PCT)
    # Perpetuity TV restated to the end of year N, divided by final-year EBITDA.
    vrow("implied", "Implied exit multiple / implied perpetuity growth",
         f"=IF({ebitda_n}>0,{{tv}}*(1+{{wacc}})^({n_ref}-{{tv_period}})/{ebitda_n},NA())",
         f"=IF({fcf_n}>0,({{tv}}/({fcf_n}*(1+{{wacc}})^({n_ref}-{_rel(P, v['tv_period'])}))*{{wacc}}-1)"
         f"/({{tv}}/({fcf_n}*(1+{{wacc}})^({n_ref}-{_rel(P, v['tv_period'])}))+1),NA())",
         NUM)
    ws_d.cell(row=v["implied"], column=P).number_format = MULT
    ws_d.cell(row=v["implied"], column=X).number_format = PCT2
    row += 1

    _section(ws_d, row, "Equity value and implied share price", 4)
    row += 1
    bridge = [("debt", "Less: total debt"), ("preferred", "Less: preferred equity"),
              ("minority", "Less: minority interest"), ("cash", "Plus: cash"),
              ("non_op", "Plus: non-operating assets")]
    for key, label in bridge:
        vrow(key, label, f"={inp[key]}", f"={inp[key]}", NUM, (GREEN, GREEN))
    vrow("equity", "Equity value", "={ev}-{debt}-{preferred}-{minority}+{cash}+{non_op}",
         "={ev}-{debt}-{preferred}-{minority}+{cash}+{non_op}", NUM, bold=True)
    vrow("basic", "Basic shares", f"={inp['basic']}", f"={inp['basic']}", SHARES, (GREEN, GREEN))
    vrow("rsus", "RSUs", f"={inp['rsus']}", f"={inp['rsus']}", SHARES, (GREEN, GREEN))
    price_row, shares_row = row, row + 1
    row += 2
    vrow("current", "Current share price",
         f"={inp['current_price']}", f"={inp['current_price']}", PRICE, (GREEN, GREEN))
    vrow("upside", "Upside / (downside)",
         f'=IF(ISNUMBER({inp["current_price"]}),{_rel(P, price_row)}/{{current}}-1,"n/a")',
         f'=IF(ISNUMBER({inp["current_price"]}),{_rel(X, price_row)}/{{current}}-1,"n/a")', PCT)
    if a.current_share_price is None:
        for c in (P, X):
            ws_d.cell(row=v["current"], column=c, value="n/a").alignment = Alignment(horizontal="right")

    # Treasury stock method. Candidate k assumes the k lowest-strike tranches are in the money and solves
    # P = (equity + option proceeds) / (basic + RSUs + options); exactly one candidate is self-consistent.
    row += 1
    _section(ws_d, row, "Treasury stock method (solves the dilution circularity without iteration)", 10)
    row += 1
    for c, h in enumerate(["Strike", "Options", "Cum. options", "Cum. proceeds",
                           "Price: perpetuity", "Valid", "Price: exit", "Valid"], start=3):
        ws_d.cell(row=row, column=c, value=h).font = Font(bold=True)
        ws_d.cell(row=row, column=c).alignment = Alignment(horizontal="right")
    row += 1
    tsm_first = row
    m = len(tranches)
    eq_p, eq_x = _abs(P, v["equity"]), _abs(X, v["equity"])
    base_sh = f"({_abs(P, v['basic'])}+{_abs(P, v['rsus'])})"
    for k in range(m + 1):
        ws_d.cell(row=row, column=2, value=("No options in the money" if k == 0 else
                                                 "Tranche 1 in the money" if k == 1 else
                                                 f"Tranches 1–{k} in the money"))
        if k == 0:
            _put(ws_d, 5, row, 0, SHARES)
            _put(ws_d, 6, row, 0, NUM)
        else:
            orow = opt_rows[k - 1]
            _put(ws_d, 3, row, f"=Inputs!$D${orow}", PRICE, GREEN)
            _put(ws_d, 4, row, f"=Inputs!$C${orow}", SHARES, GREEN)
            _put(ws_d, 5, row, f"={_rel(5, row - 1)}+{_rel(4, row)}", SHARES)
            _put(ws_d, 6, row, f"={_rel(6, row - 1)}+{_rel(4, row)}*{_rel(3, row)}", NUM)
        for price_col, eq in ((7, eq_p), (9, eq_x)):
            price = _rel(price_col, row)
            _put(ws_d, price_col, row, f"=({eq}+{_rel(6, row)})/({base_sh}+{_rel(5, row)})", PRICE)
            lower = f"{price}>{_rel(3, row)}" if k > 0 else None
            upper = f"{price}<={_rel(3, row + 1)}" if k < m else None
            conds = [c for c in (lower, upper) if c]
            test = "1" if not conds else (f"IF(AND({','.join(conds)}),1,0)")
            _put(ws_d, price_col + 1, row, f"={test}", INT)
        row += 1
    tsm_last = row - 1
    strikes = f"{_abs(3, tsm_first + 1)}:{_abs(3, tsm_last)}"
    counts = f"{_abs(4, tsm_first + 1)}:{_abs(4, tsm_last)}"
    for col, price_col in ((P, 7), (X, 9)):
        prices = f"{_abs(price_col, tsm_first)}:{_abs(price_col, tsm_last)}"
        valid = f"{_abs(price_col + 1, tsm_first)}:{_abs(price_col + 1, tsm_last)}"
        _put(ws_d, col, price_row, f"=SUMPRODUCT({prices},{valid})", PRICE, bold=True)
        price = _rel(col, price_row)
        dilution = (f"+SUMPRODUCT(({price}>{strikes})*{counts}*(1-{strikes}/{price}))" if m else "")
        _put(ws_d, col, shares_row,
             f"=IF({price}>0,{base_sh}{dilution},{base_sh})", SHARES)
    ws_d.cell(row=price_row, column=2, value="Implied share price").font = Font(bold=True)
    ws_d.cell(row=shares_row, column=2, value="Diluted shares (treasury stock method)")
    for c in range(2, 5):
        ws_d.cell(row=price_row, column=c).fill = BASE_FILL

    for key in ("tv", "pv_tv", "ev", "equity", "implied"):
        refs[f"perpetuity_{key}"] = _abs(P, v[key], "DCF")
        refs[f"exit_{key}"] = _abs(X, v[key], "DCF")
    refs["sum_pv_ufcf"] = _abs(P, v["sum_pv"], "DCF")
    refs["perpetuity_share_price"] = _abs(P, price_row, "DCF")
    refs["exit_share_price"] = _abs(X, price_row, "DCF")
    refs["perpetuity_diluted_shares"] = _abs(P, shares_row, "DCF")
    refs["exit_diluted_shares"] = _abs(X, shares_row, "DCF")

    # ---------------------------------------------------------------- Sensitivity
    _title(ws_s, f"{a.company}: sensitivity analysis",
           "Values computed by dcfmodel (Python). Rerun dcfmodel after changing inputs to refresh.")
    row = 4
    for table in tables:
        ncols = len(table.cols)
        _section(ws_s, row, table.title, 2 + ncols)
        row += 1
        ws_s.cell(row=row, column=2, value=f"{table.row_label} ↓   {table.col_label} →").font = Font(italic=True)
        col_fmt = MULT if "multiple" in table.col_label.lower() else PCT
        base_col = table.cols[len(table.cols) // 2]
        base_row = table.rows[len(table.rows) // 2]
        for j, cv in enumerate(table.cols):
            _put(ws_s, 3 + j, row, cv, col_fmt, bold=True)
        row += 1
        for i, rv in enumerate(table.rows):
            _put(ws_s, 2, row, rv, PCT2, bold=True).alignment = Alignment(horizontal="right")
            for j, value in enumerate(table.values[i]):
                cell = _put(ws_s, 3 + j, row, value if value is not None else "n/a", PRICE)
                cell.alignment = Alignment(horizontal="right")
                if rv == base_row and table.cols[j] == base_col:
                    cell.fill = BASE_FILL
                    cell.font = Font(bold=True)
            row += 1
        row += 1
    for c in range(3, 3 + max(len(t.cols) for t in tables)):
        ws_s.column_dimensions[get_column_letter(c)].width = 12

    for ws in (ws_in, ws_w, ws_d, ws_s):
        ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.active = wb.sheetnames.index("DCF")
    wb.save(path)
    return refs
