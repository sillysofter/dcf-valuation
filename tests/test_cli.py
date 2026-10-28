import json

import yaml
from openpyxl import load_workbook

from dcfmodel.cli import main
from tests.conftest import EXAMPLE


def test_cli_end_to_end(tmp_path, capsys):
    xlsx, js = tmp_path / "model.xlsx", tmp_path / "out" / "summary.json"
    assert main([str(EXAMPLE), "--excel", str(xlsx), "--json", str(js)]) == 0
    out = capsys.readouterr().out
    assert "Implied share price" in out and "WACC" in out
    report = json.loads(js.read_text())
    assert report["perpetuity_growth"]["share_price"] > 0
    assert len(report["sensitivities"]) == 2
    assert load_workbook(xlsx).sheetnames == ["Inputs", "WACC", "DCF", "Sensitivity"]


def test_cli_reports_bad_input(tmp_path, capsys):
    raw = yaml.safe_load(EXAMPLE.read_text())
    raw["terminal"]["growth_rate"] = 0.2
    bad = tmp_path / "bad.yaml"
    bad.write_text(yaml.safe_dump(raw))
    assert main([str(bad)]) == 2
    assert "below WACC" in capsys.readouterr().err
